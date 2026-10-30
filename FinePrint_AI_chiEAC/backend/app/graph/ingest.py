import csv, io, json, zipfile
from pathlib import Path
from openpyxl import load_workbook
from ..config import settings
from ..schemas import Row

def parse_file(content: bytes, filename: str) -> tuple[list[Row], list[str]]:
    if len(content) > settings.max_bytes:
        raise ValueError('File exceeds 5 MB. Split it into smaller files.')
    suffix = Path(filename).suffix.lower()
    if suffix == '.csv':
        records = list(csv.DictReader(io.StringIO(content.decode('utf-8-sig'))))
    elif suffix == '.json':
        records = json.loads(content)
        if not isinstance(records, list): raise ValueError('JSON must be an array of rows.')
    elif suffix == '.xlsx':
        with zipfile.ZipFile(io.BytesIO(content)) as z:
            if sum(i.file_size for i in z.infolist()) > 30_000_000:
                raise ValueError('Expanded spreadsheet is too large.')
        wb = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
        ws = wb.active
        if ws.max_row and ws.max_row > 401: raise ValueError('Maximum 400 rows. Split your document.')
        values = ws.iter_rows(values_only=True)
        headers = [str(x or '').strip() for x in next(values, [])]
        records = [dict(zip(headers, r)) for r in values if any(x is not None for x in r)]
        wb.close()
    else: raise ValueError('Choose CSV, JSON, or XLSX. PDF is not supported in v1.')
    if not records: raise ValueError('The file has no document rows.')
    if len(records) > settings.max_rows: raise ValueError('Maximum 400 rows. Split your document.')
    rows, warnings = [], []
    for i, record in enumerate(records):
        if isinstance(record, str): record = {'text': record}
        if not isinstance(record, dict): raise ValueError('Every row must be an object or string.')
        record = dict(record)
        if 'text' not in record:
            key = next((k for k in ['content','paragraph','clause','body'] if k in record), None)
            if key is None and len(record) == 1: key = next(iter(record))
            if key:
                record['text'] = record[key]
                warnings.append(f'Row {i+1}: used {key} as text.')
        record.setdefault('doc_id', Path(filename).stem)
        record.setdefault('section_id', f'r{i+1}')
        for k in ['heading','source_file']:
            if record.get(k) is None: record[k] = ''
        if record.get('page') == '': record['page'] = None
        rows.append(Row.model_validate(record))
    if len({r.doc_id for r in rows}) != 1: raise ValueError('Upload one document at a time; doc_id must be the same.')
    if len({r.section_id for r in rows}) != len(rows): raise ValueError('Each section_id must be unique.')
    return rows, warnings
